# Author:ADMIN
# Date:2020/3/22

# -*- coding: UTF-8 -*-
import re
from openpyxl import *
from docx import Document

class WordToExcel:
    def __init__(self):
        self.test = 0

    def get_ques_stem(self, src_doc, pattern_stem):
        dst_lst = []
        for line in src_doc.paragraphs:
            src_line = line.text
            matched_line = pattern_stem.match(src_line)
            if matched_line:
                matched_content = matched_line.group().strip()
                if matched_content:
                    dst_lst.append(matched_content)
        return dst_lst

    def get_option(self, src_doc, pattern_patition):
        tmp_lst = []
        dst_lst = []
        for line in src_doc.paragraphs:
            matched_line = pattern_patition.search(line.text)
            if matched_line:
                matched_content = matched_line.group()
                if matched_content:
                    tmp_lst.append(matched_content)

        for item in tmp_lst:
            a = re.findall(r"A(.*)B|A(.*)", item)
            if a != []:
                dst_lst.append(''.join(a[0]))
            b = re.findall(r"B(.*)C|B(.*)", item)
            if b != []:
                dst_lst.append(''.join(b[0]))
            c = re.findall(r"C(.*)D|C(.*)", item)
            if c != []:
                dst_lst.append(''.join(c[0]))

            d = re.findall(r"D(.*)", item)
            if d != []:
                dst_lst.append(''.join(d[0]))
        return dst_lst

    def get_anwser(self, src_table):
        dst_lst = []
        for table in src_table[:]:
            for i, row in enumerate(table.rows[:]):
                num = 0
                for cell in row.cells[:]:
                    c = cell.text
                    if (num % 2) == 1:
                        dst_lst.append(c)
                    num += 1
        return dst_lst

    def set_excel(self, content, rows, dst):
        wb = Workbook()
        ws = wb.active

        for row in range(0, rows):
            tmp = content[row]#依次取出每个小列表
            for col in range(0, 7):
                ws.cell(row+1, col+1, tmp[col])#从小列表中依次取出一个元素
        wb.save(dst)



